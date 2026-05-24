"""Чтение Excel-файла со списком получателей и генерация шаблона."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


PHONE_HEADERS = {"телефон", "phone", "номер", "msisdn", "tel", "mobile", "мобильный"}
NAME_HEADERS = {"имя", "name", "фио", "пользователь", "user", "fullname", "full name"}


@dataclass
class UserRow:
    name: str
    phone: str
    row_index: int  # 1-based для отчётов


def _looks_like_phone(value: object) -> bool:
    if value is None:
        return False
    return sum(ch.isdigit() for ch in str(value)) >= 7


def load_users(path: str | Path) -> List[UserRow]:
    wb = load_workbook(filename=str(path), data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = rows[0]
    name_idx: Optional[int] = None
    phone_idx: Optional[int] = None

    header_is_text = all(
        cell is None or isinstance(cell, str) for cell in header
    ) and any(isinstance(c, str) and c.strip() for c in header)

    if header_is_text:
        for i, cell in enumerate(header):
            if not isinstance(cell, str):
                continue
            v = cell.strip().lower()
            if phone_idx is None and v in PHONE_HEADERS:
                phone_idx = i
            elif name_idx is None and v in NAME_HEADERS:
                name_idx = i
        data_rows = rows[1:]
    else:
        data_rows = rows

    if phone_idx is None:
        for r in data_rows:
            for i, cell in enumerate(r):
                if _looks_like_phone(cell):
                    phone_idx = i
                    break
            if phone_idx is not None:
                break

    if phone_idx is None:
        raise ValueError("В файле не найдена колонка с номерами телефонов")

    if name_idx is None:
        for r in data_rows:
            for i, cell in enumerate(r):
                if i == phone_idx:
                    continue
                if isinstance(cell, str) and cell.strip():
                    name_idx = i
                    break
            if name_idx is not None:
                break

    out: List[UserRow] = []
    start_row = 2 if header_is_text else 1
    for offset, r in enumerate(data_rows):
        phone = r[phone_idx] if phone_idx < len(r) else None
        if phone is None or str(phone).strip() == "":
            continue
        name = ""
        if name_idx is not None and name_idx < len(r) and r[name_idx] is not None:
            name = str(r[name_idx]).strip()
        out.append(
            UserRow(
                name=name or "(без имени)",
                phone=str(phone).strip(),
                row_index=start_row + offset,
            )
        )
    return out


def write_sample(path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Пользователи"

    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill("solid", fgColor="0A84FF")
    center = Alignment(horizontal="left", vertical="center")

    ws.append(["Имя", "Телефон"])
    for col in ("A1", "B1"):
        c = ws[col]
        c.font = header_font
        c.fill = header_fill
        c.alignment = center

    ws.append(["Иван Иванов", "+7 900 123-45-67"])
    ws.append(["Мария Петрова", "79161234567"])
    ws.append(["Алексей Сидоров", "8 (905) 555-12-34"])

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.row_dimensions[1].height = 24

    wb.save(str(path))

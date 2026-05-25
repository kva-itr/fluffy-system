"""Создание звонков в VK Звонках по расписанию курса и экспорт ссылок в Excel.

Использование:

    export VK_ACCESS_TOKEN="ваш_токен"
    python vk_calls_excel.py            # создаст vk_calls.xlsx рядом со скриптом
    python vk_calls_excel.py out.xlsx   # имя выходного файла можно задать аргументом

Токен должен иметь право вызывать метод `calls.start` (VK Звонки).
Метод документирован тут: https://dev.vk.com/method/calls.start
"""

from __future__ import annotations

import os
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from typing import Optional

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


VK_API_BASE = "https://api.vk.com/method"
VK_API_VERSION = "5.199"
COURSE_TITLE = "Актуальные аспекты преподавания информатики в ОО"


@dataclass
class Session:
    num: int
    date: str           # ДД.ММ.ГГГГ
    weekday: str
    time_start: str     # ЧЧ:ММ
    time_end: str
    topic: str
    teacher: str

    @property
    def title(self) -> str:
        return f"{COURSE_TITLE}. Занятие {self.num} ({self.date} {self.time_start})"


SESSIONS: list[Session] = [
    Session(1, "25.05.2026", "пн", "10:00", "11:30",
            "Методический анализ региональных результатов ОГЭ и ЕГЭ по информатике. "
            "Анализ структуры и содержания КИМ ВПР, ОГЭ и ЕГЭ по информатике в свете "
            "требований стандартов общего основного и среднего образования.",
            "Федченко Г.М."),
    Session(2, "26.05.2026", "вт", "11:00", "12:30",
            "Профессиональные дефициты в Дальневосточном федеральном округе. "
            "Тематические блоки профориентационного характера в 10-11 классах.",
            "Матевосян А.С."),
    Session(3, "27.05.2026", "ср", "09:00", "10:30",
            "Динамическое программирование в задачах ОГЭ и ЕГЭ.",
            "Федченко Г.М."),
    Session(4, "28.05.2026", "чт", "11:30", "13:00",
            "Решение задач ОГЭ по информатике средствами LibreOffice.",
            "Иващенко Н.С."),
    Session(5, "29.05.2026", "пт", "10:00", "11:30",
            "Решение задач КЕГЭ по информатике средствами LibreOffice.",
            "Иващенко Н.С."),
    Session(6, "01.06.2026", "пн", "09:00", "10:30",
            "Решение задач различного уровня сложности методом динамического программирования.",
            "Федченко Г.М."),
    Session(7, "02.06.2026", "вт", "09:00", "10:30",
            "Решение задач ЕГЭ базового и повышенного уровней сложности "
            "с использованием электронных таблиц.",
            "Федченко Г.М."),
    Session(8, "03.06.2026", "ср", "09:00", "10:30",
            "Решение задач ЕГЭ базового и повышенного уровней сложности "
            "с помощью программирования.",
            "Федченко Г.М."),
    Session(9, "04.06.2026", "чт", "09:00", "10:30",
            "Исполнители. Моделирование работы исполнителей.",
            "Федченко Г.М."),
]


class VKAPIError(RuntimeError):
    pass


def vk_request(method: str, token: str, params: Optional[dict] = None) -> dict:
    payload = {"access_token": token, "v": VK_API_VERSION}
    if params:
        payload.update(params)
    r = requests.post(f"{VK_API_BASE}/{method}", data=payload, timeout=30)
    r.raise_for_status()
    data = r.json()
    if "error" in data:
        err = data["error"]
        raise VKAPIError(f"{err.get('error_code')}: {err.get('error_msg')}")
    return data["response"]


def create_call(token: str, session: Session) -> dict:
    """Вызывает calls.start и возвращает {'join_link': ..., 'call_id': ...}.

    Параметры метода VK API необязательны — звонок создаётся пустым,
    ссылка-приглашение приходит в join_link.
    """
    return vk_request("calls.start", token, {"name": session.title})


def build_workbook(rows: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Звонки"

    headers = ["№", "Дата", "День", "Начало", "Окончание",
               "Тема", "Преподаватель", "Ссылка для подключения",
               "ID звонка", "Статус"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        ws.append([
            row["num"], row["date"], row["weekday"],
            row["time_start"], row["time_end"],
            row["topic"], row["teacher"],
            row["join_link"], row["call_id"], row["status"],
        ])

    widths = [4, 12, 6, 8, 10, 60, 18, 55, 22, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 42
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).alignment = Alignment(
                vertical="center", wrap_text=True
            )
        link_cell = ws.cell(row=r, column=8)
        if link_cell.value and str(link_cell.value).startswith("http"):
            link_cell.hyperlink = link_cell.value
            link_cell.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "A2"
    return wb


def main() -> int:
    token = os.environ.get("VK_ACCESS_TOKEN")
    if not token:
        print("ERROR: переменная окружения VK_ACCESS_TOKEN не задана.",
              file=sys.stderr)
        return 2

    out_path = sys.argv[1] if len(sys.argv) > 1 else "vk_calls.xlsx"

    rows: list[dict] = []
    for s in SESSIONS:
        base = {
            "num": s.num, "date": s.date, "weekday": s.weekday,
            "time_start": s.time_start, "time_end": s.time_end,
            "topic": s.topic, "teacher": s.teacher,
        }
        try:
            print(f"[{s.num}/{len(SESSIONS)}] создаю звонок: {s.title}")
            resp = create_call(token, s)
            rows.append({
                **base,
                "join_link": resp.get("join_link", ""),
                "call_id": resp.get("call_id", ""),
                "status": "OK",
            })
        except (VKAPIError, requests.RequestException) as e:
            print(f"  ошибка: {e}", file=sys.stderr)
            rows.append({
                **base,
                "join_link": "",
                "call_id": "",
                "status": f"ERROR: {e}",
            })
        time.sleep(0.4)  # щадящий rate limit VK API

    wb = build_workbook(rows)
    wb.save(out_path)
    print(f"\nГотово. Файл сохранён: {os.path.abspath(out_path)}")
    print(f"Сгенерировано: {datetime.now():%Y-%m-%d %H:%M:%S}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

"""
Создаёт звонки в VK Звонках через access token и выгружает расписание
со ссылками-подключениями в Excel-файл.

Запуск:
    export VK_ACCESS_TOKEN="vk1.a.xxxxxxxx..."
    python vk_calls_export.py

Зависимости:
    pip install requests openpyxl
"""

from __future__ import annotations

import os
import sys
import time
from dataclasses import dataclass

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


VK_API = "https://api.vk.com/method"
VK_API_VERSION = "5.199"


@dataclass
class Session:
    num: int
    date: str
    time: str
    topic: str
    teacher: str


SESSIONS: list[Session] = [
    Session(1, "25.05, пн", "10.00-11.30",
            "Методический анализ региональных результатов ОГЭ и ЕГЭ по информатике. "
            "Анализ структуры и содержания КИМ ВПР, ОГЭ и ЕГЭ по информатике "
            "в свете требований стандартов общего основного и среднего образования.",
            "Федченко Г.М."),
    Session(2, "26.05, вт", "11.00-12.30",
            "Профессиональные дефициты в Дальневосточном федеральном округе. "
            "Тематические блоки профориентационного характера в 10-11 классах.",
            "Матевосян А.С."),
    Session(3, "27.05, ср", "09.00-10.30",
            "Динамическое программирование в задачах ОГЭ и ЕГЭ.",
            "Федченко Г.М."),
    Session(4, "28.05, чт", "11.30-13.00",
            "Решение задач ОГЭ по информатике средствами LibreOffice.",
            "Иващенко Н.С."),
    Session(5, "29.05, пт", "10.00-11.30",
            "Решение задач КЕГЭ по информатике средствами LibreOffice.",
            "Иващенко Н.С."),
    Session(6, "01.06, пн", "09.00-10.30",
            "Решение задач различного уровня сложности методом динамического программирования.",
            "Федченко Г.М."),
    Session(7, "02.06, вт", "09.00-10.30",
            "Решение задач ЕГЭ базового и повышенного уровней сложности "
            "с использованием электронных таблиц.",
            "Федченко Г.М."),
    Session(8, "03.06, ср", "09.00-10.30",
            "Решение задач ЕГЭ базового и повышенного уровней сложности "
            "с помощью программирования.",
            "Федченко Г.М."),
    Session(9, "04.06, чт", "09.00-10.30",
            "Исполнители. Моделирование работы исполнителей.",
            "Федченко Г.М."),
]


def create_vk_call(access_token: str) -> str:
    """Создаёт звонок в VK и возвращает join_link."""
    resp = requests.get(
        f"{VK_API}/calls.start",
        params={"access_token": access_token, "v": VK_API_VERSION},
        timeout=30,
    )
    resp.raise_for_status()
    data = resp.json()
    if "error" in data:
        err = data["error"]
        raise RuntimeError(
            f"VK API error {err.get('error_code')}: {err.get('error_msg')}"
        )
    return data["response"]["join_link"]


def export_to_excel(rows: list[tuple[Session, str]], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Звонки"

    title = ("АКТУАЛЬНЫЕ АСПЕКТЫ ПРЕПОДАВАНИЯ ИНФОРМАТИКИ "
             "В ОБЩЕОБРАЗОВАТЕЛЬНОЙ ОРГАНИЗАЦИИ")
    period = "25 мая – 7 июня 2026 г., время МСК"

    ws.merge_cells("A1:F1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("A2:F2")
    ws["A2"] = period
    ws["A2"].font = Font(italic=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["№", "Дата", "Время", "Тема", "Преподаватель", "Ссылка для подключения"]
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [5, 14, 14, 60, 18, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for i, (s, link) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=s.num).alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=2, value=s.date)
        ws.cell(row=i, column=3, value=s.time)
        ws.cell(row=i, column=4, value=s.topic).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=5, value=s.teacher).alignment = Alignment(wrap_text=True, vertical="top")
        link_cell = ws.cell(row=i, column=6, value=link)
        if link.startswith("http"):
            link_cell.hyperlink = link
            link_cell.font = Font(color="0563C1", underline="single")
        link_cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


def main() -> int:
    token = os.environ.get("VK_ACCESS_TOKEN")
    if not token:
        print("Ошибка: переменная окружения VK_ACCESS_TOKEN не задана.", file=sys.stderr)
        print('Пример: export VK_ACCESS_TOKEN="vk1.a.xxxx..."', file=sys.stderr)
        return 1

    out_path = sys.argv[1] if len(sys.argv) > 1 else "vk_calls_schedule.xlsx"

    rows: list[tuple[Session, str]] = []
    for s in SESSIONS:
        print(f"[{s.num}/{len(SESSIONS)}] Создаю звонок для {s.date} {s.time}…")
        try:
            link = create_vk_call(token)
        except Exception as exc:
            print(f"  Не удалось создать звонок: {exc}", file=sys.stderr)
            link = f"ERROR: {exc}"
        else:
            print(f"  Ссылка: {link}")
        rows.append((s, link))
        time.sleep(0.4)  # лёгкий троттлинг под лимиты VK API

    export_to_excel(rows, out_path)
    print(f"\nГотово. Файл сохранён: {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
